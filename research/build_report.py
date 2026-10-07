"""
Builds the Word (.docx) research report from the outputs of the four
analysis scripts. Every number, table and conclusion in the text is read
from those outputs at build time -- nothing is typed in by hand -- so after
re-running the experiments (e.g. the full VM run), re-running this script
regenerates a report that is consistent with the new data.

Reads:
  output/metrics_comparison/   run_metrics_comparison.py
  output/hparam_sweeps/        run_hparam_sweeps.py
  output/burst_comparisons/    plot_burst_comparison.py
  output/statistical_analysis/ run_statistical_analysis.py
plus the dataset's trace files (for the network statistics).

If the sample sizes are too small to support conclusions, a visible
"preliminary results" notice is inserted at the top of the report.

Usage:
    python -m research.build_report
"""
import argparse
import datetime
import importlib.metadata
import json
import platform
import re
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

from research import dataset_utils, concealment, safe_io
from research.run_metrics_comparison import ALL_METRICS

OUTPUT_DIR = Path(__file__).resolve().parent / "output"
DEFAULT_OUT = OUTPUT_DIR / "report" / "PLC_LPC_vs_PARCnet_Report.docx"

AUTHOR = "Omar Ali Elgabry"
TITLE = ("Hybrid Neural versus Purely Autoregressive Packet Loss Concealment "
         "for Networked Music: An Evaluation on Real Loss Traces")

MIN_FILES_FOR_CONCLUSIONS = 30
MIN_SWEEP_FILES = 10
CONTENT_WIDTH_CM = 16.0

METRIC_DESC = {
    "snr": "time-domain signal-to-noise ratio",
    "mrstft": "multi-resolution STFT distance",
    "plcmos": "PLCMOS",
    "visqol": "ViSQOL MOS-LQO",
}
SWEEP_LABELS = {
    "ar_order": ("AR order p", "ar_order"),
    "context_length_packets": ("Context length (packets)", "context_dim_packets"),
    "diagonal_load": ("Diagonal loading λ", "diagonal_load"),
    "extra_dim": ("Crossfade length E (samples)", "extra_dim"),
}

REFERENCES = [
    ("mezza2024", 'A. I. Mezza, M. Amerena, A. Bernardini, and A. Sarti, "Hybrid packet loss concealment for '
                  'real-time networked music applications," IEEE Open Journal of Signal Processing, vol. 5, '
                  'pp. 266-273, 2024, doi: 10.1109/OJSP.2023.3343318.'),
    ("challenge2024", 'A. I. Mezza et al., "Report on the IEEE-IS² 2024 Music Packet Loss Concealment Challenge," '
                      'arXiv:2409.18564, 2024.'),
    ("perkins1998", 'C. Perkins, O. Hodson, and V. Hardman, "A survey of packet loss recovery techniques for '
                    'streaming audio," IEEE Network, vol. 12, no. 5, pp. 40-48, 1998.'),
    ("makhoul1975", 'J. Makhoul, "Linear prediction: A tutorial review," Proceedings of the IEEE, vol. 63, no. 4, '
                    'pp. 561-580, 1975.'),
    ("lostanlen2016", 'V. Lostanlen and C. E. Cella, "Deep convolutional networks on the pitch spiral for musical '
                      'instrument recognition," in Proc. ISMIR, 2016.'),
    ("diener2022", 'L. Diener, S. Sootla, S. Branets, A. Saabas, R. Aichner, and R. Cutler, "INTERSPEECH 2022 '
                   'Audio Deep Packet Loss Concealment Challenge," in Proc. Interspeech, 2022.'),
    ("yamamoto2020", 'R. Yamamoto, E. Song, and J.-M. Kim, "Parallel WaveGAN: A fast waveform generation model '
                     'based on generative adversarial networks with multi-resolution spectrogram," in Proc. '
                     'IEEE ICASSP, 2020.'),
    ("diener2023", 'L. Diener, M. Purin, S. Sootla, A. Saabas, R. Aichner, and R. Cutler, "PLCMOS - a data-driven '
                   'non-intrusive metric for the evaluation of packet loss concealment algorithms," in Proc. '
                   'Interspeech, 2023.'),
    ("chinen2020", 'M. Chinen, F. S. C. Lim, J. Skoglund, N. Gureev, F. O\'Gorman, and A. Hines, "ViSQOL v3: An '
                   'open source production ready objective speech and audio metric," in Proc. QoMEX, 2020.'),
    ("gray1974", 'A. H. Gray and J. D. Markel, "A spectral-flatness measure for studying the autocorrelation '
                 'method of linear prediction of speech analysis," IEEE Trans. Acoustics, Speech, and Signal '
                 'Processing, vol. 22, no. 3, pp. 207-217, 1974.'),
    ("wilcoxon1945", 'F. Wilcoxon, "Individual comparisons by ranking methods," Biometrics Bulletin, vol. 1, '
                     'no. 6, pp. 80-83, 1945.'),
    ("holm1979", 'S. Holm, "A simple sequentially rejective multiple test procedure," Scandinavian Journal of '
                 'Statistics, vol. 6, no. 2, pp. 65-70, 1979.'),
    ("kerby2014", 'D. S. Kerby, "The simple difference formula: An approach to teaching nonparametric '
                  'correlation," Comprehensive Psychology, vol. 3, 2014.'),
    ("efron1993", 'B. Efron and R. J. Tibshirani, An Introduction to the Bootstrap. Chapman & Hall/CRC, 1993.'),
    ("itu1534", 'ITU-R Recommendation BS.1534-3, "Method for the subjective assessment of intermediate quality '
                'level of audio systems (MUSHRA)," 2015.'),
]
REF_NUM = {key: i + 1 for i, (key, _) in enumerate(REFERENCES)}


def cite(*keys) -> str:
    return "[" + ", ".join(str(REF_NUM[k]) for k in keys) + "]"


# --------------------------------------------------------------------------
# Number formatting
# --------------------------------------------------------------------------
def f(v, d=2) -> str:
    return "n/a" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:.{d}f}"


def fs(v, d=2) -> str:
    return "n/a" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:+.{d}f}"


def fp(p) -> str:
    if p is None or np.isnan(p):
        return "n/a"
    return "< 0.001" if p < 0.001 else f"{p:.3f}"


def fg(v) -> str:
    return f"{v:g}"


def join_words(items: list[str]) -> str:
    items = list(items)
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " and " + items[-1]


# --------------------------------------------------------------------------
# Document helpers
# --------------------------------------------------------------------------
def shade(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def add_field(run, instruction: str, placeholder: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, text, end):
        run._r.append(el)


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self._setup()

    def _setup(self):
        section = self.doc.sections[0]
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.5))

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for level, size in [(1, 15), (2, 12.5), (3, 11.5)]:
            h = styles[f"Heading {level}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            h.paragraph_format.space_after = Pt(4)
        caption = styles["Caption"]
        caption.font.size = Pt(9)
        caption.font.italic = False
        caption.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

        # Ask Word to refresh fields (table of contents, page numbers) on open.
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        self.doc.settings.element.append(update)

        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_p.add_run()
        run.font.size = Pt(9)
        add_field(run, "PAGE", "1")

    # -- text --------------------------------------------------------------
    def heading(self, text: str, level: int = 1):
        self.doc.add_heading(text, level=level)

    def para(self, *parts, align=None, size=None, space_after=None, style=None):
        """parts: plain strings, or (text, flags) tuples with flags from 'b', 'i', 'sub', 'sup'."""
        p = self.doc.add_paragraph(style=style)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for part in parts:
            text, flags = (part, "") if isinstance(part, str) else part
            run = p.add_run(text)
            run.bold = "b" in flags.replace("sub", "")
            run.italic = "i" in flags
            run.font.subscript = "sub" in flags
            run.font.superscript = "sup" in flags
            if size:
                run.font.size = Pt(size)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def body(self, *parts):
        return self.para(*parts, align="justify")

    def bullet(self, *parts):
        return self.para(*parts, style="List Bullet")

    def numbered(self, *parts):
        return self.para(*parts, style="List Number")

    def equation(self, text: str) -> int:
        self.eq_n += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run(f"\t{text}\t({self.eq_n})")
        run.font.name = "Cambria Math"
        run.italic = True
        return self.eq_n

    def callout(self, title: str, text: str, fill="FFF4E5"):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        shade(cell, fill)
        p = cell.paragraphs[0]
        r = p.add_run(title + " ")
        r.bold = True
        p.add_run(text)
        for run in p.runs:
            run.font.size = Pt(10)
        self.doc.add_paragraph()

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        p = self.doc.add_paragraph()
        add_field(p.add_run(), 'TOC \\o "1-2" \\h \\z \\u',
                  "Right-click here and choose 'Update Field' to build the table of contents.")

    # -- figures & tables ----------------------------------------------------
    def next_fig(self) -> int:
        return self.fig_n + 1

    def next_tab(self) -> int:
        return self.tab_n + 1

    def figure(self, path: Path, caption: str, width_cm: float = CONTENT_WIDTH_CM) -> int | None:
        if not path or not Path(path).exists():
            self.para((f"[Figure not available: {Path(path).name if path else 'missing'} -- "
                       f"run the script that produces it]", "i"), size=9)
            return None
        self.fig_n += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.doc.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.add_run(f"Figure {self.fig_n}. ").bold = True
        cap.add_run(caption)
        cap.paragraph_format.space_after = Pt(12)
        return self.fig_n

    def table(self, header: list[str], rows: list[list], caption: str, col_widths_cm=None,
              font_size: float = 8.5, notes: str | None = None) -> int:
        self.tab_n += 1
        cap = self.doc.add_paragraph(style="Caption")
        cap.paragraph_format.keep_with_next = True
        cap.add_run(f"Table {self.tab_n}. ").bold = True
        cap.add_run(caption)

        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, text in enumerate(header):
            cell = table.rows[0].cells[j]
            shade(cell, "E5E7EB")
            run = cell.paragraphs[0].add_run(str(text))
            run.bold = True
            run.font.size = Pt(font_size)
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                run = table.rows[i].cells[j].paragraphs[0].add_run(str(value))
                run.font.size = Pt(font_size)
        if col_widths_cm:
            for row in table.rows:
                for cell, w in zip(row.cells, col_widths_cm):
                    cell.width = Cm(w)
        if notes:
            n = self.doc.add_paragraph()
            r = n.add_run(notes)
            r.italic = True
            r.font.size = Pt(8)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.tab_n


# --------------------------------------------------------------------------
# Data loading
# --------------------------------------------------------------------------
def read_csv(path: Path) -> pd.DataFrame | None:
    return pd.read_csv(path) if path.exists() else None


def network_statistics() -> dict:
    traces_dir = dataset_utils.EXAMPLE_TEST_SET_DIR / "traces"
    stems = dataset_utils.list_available_stems()
    lengths, per_file_plr, total_packets, total_lost = [], [], 0, 0
    for stem in stems:
        trace = np.loadtxt(traces_dir / f"{stem}.txt", dtype=int).flatten()
        if len(trace) == 0:
            continue
        total_packets += len(trace)
        total_lost += int(trace.sum())
        per_file_plr.append(100 * trace.mean())
        lengths.extend(length for _, length in dataset_utils.find_bursts(trace))
    lengths = np.array(lengths)
    per_file_plr = np.array(per_file_plr)
    return {
        "n_files": len(stems),
        "total_packets": total_packets,
        "total_lost": total_lost,
        "plr": 100 * total_lost / total_packets,
        "n_bursts": len(lengths),
        "mean_burst": float(lengths.mean()),
        "max_burst": int(lengths.max()),
        "pct_single": 100 * float((lengths == 1).mean()),
        "pct_le2": 100 * float((lengths <= 2).mean()),
        "n_ge8": int((lengths >= 8).sum()),
        "plr_median": float(np.median(per_file_plr)),
        "plr_min": float(per_file_plr.min()),
        "plr_max": float(per_file_plr.max()),
        "n_zero_loss": int((per_file_plr == 0).sum()),
    }


def software_versions() -> list[list[str]]:
    rows = [["Python", platform.python_version()]]
    for pkg in ["numpy", "scipy", "pandas", "librosa", "torch", "pytorch-lightning", "numba",
                "speechmos", "visqol-python", "matplotlib", "seaborn", "python-docx"]:
        try:
            rows.append([pkg, importlib.metadata.version(pkg)])
        except importlib.metadata.PackageNotFoundError:
            pass
    return rows


def oriented(df: pd.DataFrame, key: str, lower_is_better: bool) -> pd.Series:
    d = df[f"{key}_nn"] - df[f"{key}_lpc"]
    return -d if lower_is_better else d


def burst_figure_sets(directory: Path, prefix_filter: str = "") -> list[tuple[int, str, dict]]:
    """[(burst_length, tag, {'compare': path, 'diff': path, 'crossfade': path})] sorted by length."""
    sets = []
    for compare in sorted(directory.glob(f"compare_{prefix_filter}*.png")):
        m = re.match(r"compare_(.+)_(\d+)pkt\.png$", compare.name)
        if not m:
            continue
        tag, length = m.group(1), int(m.group(2))
        paths = {kind: directory / f"{kind}_{tag}_{length}pkt.png" for kind in ("compare", "diff", "crossfade")}
        sets.append((length, tag, paths))
    return sorted(sets, key=lambda s: (s[0], s[1]))


# --------------------------------------------------------------------------
# Report sections
# --------------------------------------------------------------------------
class Context:
    """All inputs loaded once, so every section reads the same numbers."""

    def __init__(self, output_dir: Path):
        self.out = output_dir
        metrics_dir = output_dir / "metrics_comparison"
        stats_dir = output_dir / "statistical_analysis"
        self.per_file = read_csv(metrics_dir / "per_file_results.csv")
        self.per_burst = read_csv(metrics_dir / "per_burst_results.csv")
        info_path = metrics_dir / "run_info.json"
        self.run_info = json.loads(info_path.read_text()) if info_path.exists() else {}
        self.hparams = self.run_info.get("hparams", dict(ar_order=256, diagonal_load=0.001,
                                                         context_dim_packets=8, extra_dim=256))
        self.tests = read_csv(stats_dir / "significance_tests.csv")
        self.burst_q = read_csv(stats_dir / "burst_length_quality.csv")
        self.sweep_summary = read_csv(stats_dir / "sweep_summary.csv")
        self.failures = read_csv(stats_dir / "failure_cases" / "failure_cases.csv")
        self.gap_mse = read_csv(output_dir / "burst_comparisons" / "gap_mse_summary.csv")
        self.sweeps = {}
        for path in sorted((output_dir / "hparam_sweeps").glob("sweep_*.csv")):
            df = pd.read_csv(path)
            self.sweeps[df["param"].iloc[0]] = (df, path.with_suffix(".png"))
        self.net = network_statistics()

        self.n_files = len(self.per_file) if self.per_file is not None else 0
        self.sr = int(self.per_file["sample_rate"].iloc[0]) if (
            self.per_file is not None and "sample_rate" in self.per_file) else 44100
        self.packet_ms = concealment.PACKET_SIZE / self.sr * 1000
        self.has_timing = self.per_file is not None and "ms_per_lost_packet_nn" in self.per_file
        self.sweep_files = {p: df["file"].nunique() for p, (df, _) in self.sweeps.items()}
        self.preliminary = (self.n_files < MIN_FILES_FOR_CONCLUSIONS or
                            any(n < MIN_SWEEP_FILES for n in self.sweep_files.values()))

    def test_row(self, key: str):
        label = next(l for k, l, _ in ALL_METRICS if k == key)
        return self.tests[self.tests["metric"] == label].iloc[0]

    def rtf(self, method: str) -> np.ndarray:
        return (self.per_file[f"ms_per_lost_packet_{method}"] / self.packet_ms).dropna().to_numpy()


def significance_sentence(ctx: Context) -> str:
    if ctx.tests is None:
        return ""
    groups = {"NN better": [], "LPC better": [], "no significant difference": []}
    for _, row in ctx.tests.iterrows():
        groups[row["verdict"]].append(row["metric"])
    parts = []
    if groups["NN better"]:
        parts.append(f"significantly better with the neural residual on {join_words(groups['NN better'])}")
    if groups["LPC better"]:
        parts.append(f"significantly better with AR alone on {join_words(groups['LPC better'])}")
    if groups["no significant difference"]:
        parts.append(f"not significantly different on {join_words(groups['no significant difference'])}")
    return f"After Holm correction (α = 0.05, n = {ctx.n_files} files), quality was " + "; ".join(parts) + "."


def realtime_sentence(ctx: Context) -> str:
    if not ctx.has_timing:
        return ""
    lpc, nn = np.median(ctx.rtf("lpc")), np.median(ctx.rtf("nn"))
    return (f"On the evaluation CPU, concealing one {ctx.packet_ms:.1f} ms packet took a median of "
            f"{np.median(ctx.per_file['ms_per_lost_packet_lpc'].dropna()):.2f} ms with AR alone "
            f"(real-time factor {lpc:.2f}) and {np.median(ctx.per_file['ms_per_lost_packet_nn'].dropna()):.1f} ms "
            f"with the hybrid model (real-time factor {nn:.2f}), i.e. the hybrid model "
            f"{'cannot' if nn > 1 else 'can'} run in real time on this hardware.")


def burst_sentence(ctx: Context) -> str:
    if ctx.burst_q is None or ctx.burst_q.empty:
        return ""
    first, last = ctx.burst_q.iloc[0], ctx.burst_q.iloc[-1]
    sig_pos = [r["burst_length_bin"] for _, r in ctx.burst_q.iterrows() if r["nn_advantage_ci_low"] > 0]
    sig_neg = [r["burst_length_bin"] for _, r in ctx.burst_q.iterrows() if r["nn_advantage_ci_high"] < 0]
    s = (f"Gap SNR decreased from {f(first['gap_snr_lpc_mean'])} dB (AR) / {f(first['gap_snr_nn_mean'])} dB "
         f"(hybrid) for single-packet losses to {f(last['gap_snr_lpc_mean'])} / {f(last['gap_snr_nn_mean'])} dB "
         f"for bursts of {last['burst_length_bin']} packets")
    if not sig_pos and not sig_neg:
        s += ", and the hybrid model's gap-SNR advantage was not distinguishable from zero at any burst length."
    else:
        bits = []
        if sig_pos:
            bits.append(f"positive for bursts of {join_words(sig_pos)} packets")
        if sig_neg:
            bits.append(f"negative for bursts of {join_words(sig_neg)} packets")
        s += f"; the hybrid model's gap-SNR advantage was {join_words(bits)} (95% CI excluding zero)."
    return s


def section_front(r: Report, ctx: Context):
    r.para((TITLE, "b"), align="center", size=17, space_after=10)
    r.para(AUTHOR, align="center", size=12, space_after=2)
    r.para(datetime.date.today().strftime("%B %Y"), align="center", size=10, space_after=16)

    if ctx.preliminary:
        thin = [f"{SWEEP_LABELS.get(p, (p,))[0]} ({n})" for p, n in ctx.sweep_files.items() if n < MIN_SWEEP_FILES]
        msg = (f"This version of the report was generated from small test runs: the main comparison uses "
               f"{ctx.n_files} files (at least {MIN_FILES_FOR_CONCLUSIONS} are needed for reliable statistics)")
        if thin:
            msg += f", and the hyperparameter sweeps were averaged over too few files ({join_words(thin)})"
        msg += (". All numbers and conclusions below will change when the full experiments are run; "
                "regenerate this report afterwards with `python -m research.build_report`.")
        r.callout("PRELIMINARY RESULTS.", msg)

    r.para(("Abstract", "b"), size=11, space_after=2)
    abstract = (
        f"Packet loss concealment (PLC) is indispensable in networked music performance, where retransmission "
        f"is impossible within the latency budget and a lost {concealment.PACKET_SIZE}-sample packet must be "
        f"replaced causally, from past audio only. This report compares two causal concealers that share the "
        f"same autoregressive (AR) linear predictor: AR prediction alone, and PARCnet, a hybrid model that adds "
        f"a convolutional neural network's estimate of the AR residual {cite('mezza2024')}. Both were evaluated on "
        f"{ctx.n_files} music recordings from the IEEE-IS² 2024 Music PLC Challenge test set with their real "
        f"packet-loss traces ({ctx.net['plr']:.1f}% of packets lost overall), using two physical metrics "
        f"(SNR and multi-resolution STFT distance) and two perceptual metrics (PLCMOS and ViSQOL), paired "
        f"non-parametric statistics, a per-burst analysis over every real loss event, one-at-a-time "
        f"hyperparameter sweeps and a real-time cost analysis. ")
    abstract += " ".join(s for s in [significance_sentence(ctx), burst_sentence(ctx), realtime_sentence(ctx)] if s)
    r.body(abstract)
    r.para(("Keywords: ", "b"), ("packet loss concealment; networked music performance; linear prediction; "
                                 "hybrid neural networks; PARCnet; objective audio quality", "i"), size=10)
    r.page_break()
    r.para(("Contents", "b"), size=13)
    r.toc()
    r.page_break()


def section_introduction(r: Report, ctx: Context):
    r.heading("1  Introduction", 1)
    r.body(f"In networked music performance (NMP), musicians at different locations play together over the "
           f"internet. Audio is sent in small packets, and every network hop can drop or delay one. Unlike "
           f"speech calls or streaming, NMP leaves no room for retransmission or large jitter buffers: the "
           f"end-to-end latency must stay within a few tens of milliseconds for ensemble playing to remain "
           f"possible. A lost packet must therefore be replaced, on the fly, by a plausible estimate computed "
           f"only from audio that has already arrived. This task is packet loss concealment (PLC) "
           f"{cite('perkins1998')}.")
    r.body(f"A classic and computationally cheap approach is to fit an autoregressive (AR) model, i.e. a "
           f"linear predictor, to the most recent audio and extrapolate it across the gap {cite('makhoul1975')}. "
           f"Linear prediction captures the stationary, resonant part of a musical signal well, but it cannot "
           f"anticipate changes such as note onsets, and its extrapolation decays during long gaps. PARCnet "
           f"{cite('mezza2024')} keeps the linear predictor and adds a lightweight convolutional network that "
           f"is trained to estimate what the predictor gets wrong, its residual. A modified PARCnet is the "
           f"official baseline of the IEEE-IS² 2024 Music PLC Challenge {cite('challenge2024')}, whose "
           f"settings ({concealment.PACKET_SIZE}-sample packets at 44.1 kHz, i.e. {ctx.packet_ms:.1f} ms, and "
           f"strictly causal processing) are adopted here.")
    r.body("Because the hybrid model contains the AR predictor as a component, comparing the two isolates "
           "exactly one question: what does the neural residual add? This report answers it through four "
           "research questions:")
    r.numbered(("RQ1 (quality). ", "b"), "Does the neural residual improve concealment quality over AR "
               "prediction alone, as measured by physical and by perceptual metrics, on real loss traces?")
    r.numbered(("RQ2 (conditions). ", "b"), "How does any advantage depend on the length of the loss burst "
               "and on the complexity of the music signal?")
    r.numbered(("RQ3 (sensitivity). ", "b"), "How sensitive are both methods to their main hyperparameters "
               "(AR order, context length, diagonal loading, crossfade length)?")
    r.numbered(("RQ4 (cost). ", "b"), "Can each method run in real time on a general-purpose CPU?")
    r.body("The main contributions are: (i) two independently implemented concealers that are verified to be "
           "bit-identical at the AR level, so that every measured difference is attributable to the neural "
           "network; (ii) an evaluation on whole files with their real packet traces rather than synthetic "
           "losses; (iii) paired statistical testing with multiple-comparison correction and effect sizes; "
           "(iv) a per-burst analysis over every real loss event in the evaluated files; and (v) hyperparameter "
           "sensitivity and real-time cost analyses.")


def section_methods(r: Report, ctx: Context):
    hp = ctx.hparams
    P, E = concealment.PACKET_SIZE, hp["extra_dim"]
    C = hp["context_dim_packets"] * P
    r.heading("2  Methods", 1)

    r.heading("2.1  Problem formulation", 2)
    r.body(f"A mono music signal x[n], sampled at f", ("s", "sub"), f" = {ctx.sr / 1000:.1f} kHz, is divided into "
           f"packets of P = {P} samples ({ctx.packet_ms:.1f} ms). A binary packet trace l", ("k", "sub"),
           " ∈ {0, 1} marks packet k as received (0) or lost (1); lost packets arrive as zeros. A concealer "
           "must output an estimate of each lost packet using only samples before it (received or previously "
           "concealed). Because concealed samples become context for the next prediction, both methods are "
           "causal and recursive: a burst of N lost packets is concealed by N successive predictions, each "
           "conditioned on the previous ones.")

    r.heading("2.2  AR (linear-prediction) concealment", 2)
    r.body(f"For every lost packet, an AR model of order p is fitted to the most recent C = "
           f"{hp['context_dim_packets']} packets ({C} samples, {C / ctx.sr * 1000:.0f} ms) of the output signal:")
    r.equation("x̂[n] = Σᵢ₌₁ᵖ aᵢ · x[n − i]")
    r.body("The coefficients a = (a₁, …, aₚ) are obtained with the autocorrelation method and solved with the "
           "Levinson-Durbin recursion, after adding a diagonal loading term λ to the zero-lag autocorrelation "
           "(white-noise compensation), which improves the conditioning of the system "
           f"{cite('makhoul1975')}:")
    r.equation("(R + λI) a = r,     Rᵢⱼ = rₓₓ(|i − j|)")
    r.body(f"The model is then run forward recursively, feeding each predicted sample back as input, to produce "
           f"P + E = {P + E} samples: the lost packet plus E = {E} extra samples ({E / ctx.sr * 1000:.1f} ms). "
           f"The extra samples are used for a linear crossfade: the tail of the prediction fades out while the "
           f"next packet fades in, so that no discontinuity is heard where concealment hands back to real audio. "
           f"Within a burst, the next prediction likewise fades in over the previous prediction's tail. The "
           f"default hyperparameters (Table {r.next_tab()}) are those of the challenge baseline.")
    r.table(["Hyperparameter", "Symbol", "Default", "Meaning"], [
        ["AR order", "p", hp["ar_order"], "number of past samples in the linear predictor"],
        ["Context length", "C", f"{hp['context_dim_packets']} packets ({C} samples)",
         "audio used to fit the AR model and as NN input"],
        ["Diagonal loading", "λ", fg(hp["diagonal_load"]), "regularization of the autocorrelation matrix"],
        ["Crossfade length", "E", f"{E} samples ({E / ctx.sr * 1000:.1f} ms)",
         "extra predicted samples used for crossfading"],
        ["NN fade-in", "-", f"{concealment.NN_FADE_DIM} samples", "fade-in of the NN contribution (hybrid only)"],
    ], "Default hyperparameters (IEEE-IS² 2024 challenge baseline configuration).",
        col_widths_cm=[3.3, 1.5, 4.2, 7.0])

    r.heading("2.3  Hybrid AR + neural concealment (PARCnet)", 2)
    r.body("The hybrid model computes the same AR prediction and adds the output of a neural network f", ("θ", "i"),
           " that receives the same past context c:")
    r.equation("ŝ = ŝ_AR + f_θ(c)")
    r.body(f"The network is a lightweight, fully convolutional 1-D encoder-decoder ('lite' configuration, 8 base "
           f"channels): four dilated residual blocks with ×2 downsampling, a bottleneck of six gated linear unit "
           f"(GLU) blocks with dilations 1 to 32, four upsampling dilated residual blocks, and a tanh output. It "
           f"was trained by the challenge organizers on Medley-solos-DB {cite('lostanlen2016')} to minimize")
    r.equation("L = 100 · ‖s − ŝ‖₁ + ½ (L_SC + L_mag)")
    r.body(f"where L_SC and L_mag are the spectral-convergence and log-magnitude terms of a multi-resolution STFT "
           f"loss {cite('yamamoto2020')}. Because the target of the network is s − ŝ_AR, it learns the residual "
           f"of an AR predictor with the default configuration. The pretrained challenge checkpoint is used "
           f"unchanged; it was not retrained for any of the swept configurations (Section 2.8). The NN output is "
           f"faded in over its first {concealment.NN_FADE_DIM} samples, and the same crossfade as in Section 2.2 "
           f"is applied to the combined prediction.")

    r.heading("2.4  Implementation and verification", 2)
    r.body("The two concealers are separate code paths. The AR concealer is a standalone implementation with no "
           "dependency on the neural network or its framework; the hybrid concealer is the unmodified challenge "
           "implementation with the network always enabled. With identical hyperparameters, the standalone AR "
           "concealer was verified to produce output identical to the hybrid model's internal AR branch "
           "(mean squared error 0.0). Every difference reported below can therefore be attributed to the neural "
           "residual alone, and not to implementation details.")

    r.heading("2.5  Dataset and loss traces", 2)
    n = ctx.net
    r.body(f"Experiments use the example test set released with the IEEE-IS² 2024 Music PLC Challenge "
           f"{cite('challenge2024')}: {n['n_files']} single-channel recordings of solo instrumental music at "
           f"44.1 kHz, each provided as a clean file, a lossy file and the packet trace that produced it. The "
           f"traces are real network traces, repurposed by the challenge from the INTERSPEECH 2022 Audio Deep "
           f"PLC Challenge {cite('diener2022')}, so no synthetic loss model is involved. Table {r.next_tab()} "
           f"summarizes the loss statistics over all {n['n_files']} traces.")
    ctx.net_table = r.table(["Statistic", "Value"], [
        ["Recordings (clean / lossy / trace triplets)", n["n_files"]],
        ["Total packets", f"{n['total_packets']:,}"],
        ["Lost packets (overall packet loss rate)", f"{n['total_lost']:,} ({n['plr']:.2f}%)"],
        ["Per-file loss rate: median (min-max)", f"{n['plr_median']:.2f}% ({n['plr_min']:.2f}-{n['plr_max']:.2f}%)"],
        ["Files without any loss", n["n_zero_loss"]],
        ["Loss bursts (runs of consecutive lost packets)", f"{n['n_bursts']:,}"],
        ["Mean / maximum burst length", f"{n['mean_burst']:.2f} / {n['max_burst']} packets"],
        ["Bursts of 1 packet / of at most 2 packets", f"{n['pct_single']:.1f}% / {n['pct_le2']:.1f}%"],
        ["Bursts of 8 packets or more", n["n_ge8"]],
    ], "Packet-loss statistics of the real traces of the test set.", col_widths_cm=[9.5, 6.5])
    seed = ctx.run_info.get("seed", "n/a")
    r.body(f"The main comparison uses {ctx.n_files} of these files (random seed {seed}). Files are drawn by a "
           f"seeded shuffle; when fewer files are requested than the candidate pool holds, the sample is "
           f"stratified across five quantile bins of spectral flatness, so that tonal and noise-like material "
           f"are both represented. Every file is concealed whole, against its own real trace.")

    r.heading("2.6  Objective quality metrics", 2)
    r.body("Two families of metrics are used, because a concealment can be close to the original waveform "
           "without sounding natural, and vice versa:")
    r.bullet(("SNR (dB, higher is better). ", "b"), "Time-domain signal-to-noise ratio between the clean and the "
             "concealed signal. It rewards an exact waveform match, including phase.")
    r.bullet(("Multi-resolution STFT distance (lower is better). ", "b"), "Spectral convergence plus mean "
             "absolute log-magnitude error, averaged over FFT sizes of 512, 1024 and 2048 "
             f"{cite('yamamoto2020')}. It compares magnitude spectra and ignores phase.")
    r.bullet(("PLCMOS (≈1-5, higher is better). ", "b"), f"A non-intrusive neural estimate of the mean opinion "
             f"score, designed specifically for packet loss concealment {cite('diener2023')}; computed at 16 kHz.")
    r.bullet(("ViSQOL MOS-LQO (≈1-5, higher is better). ", "b"), f"An intrusive perceptual similarity model, "
             f"used in its full-band audio mode at 48 kHz {cite('chinen2020')}.")
    r.body("PLCMOS and ViSQOL need seconds of audio and are therefore computed on whole files. To study individual "
           "loss events, the SNR is also computed inside each gap G (the lost samples only):")
    r.equation("SNR_gap = 10 · log₁₀( Σₙ∊G x[n]² / Σₙ∊G (x[n] − x̂[n])² )")
    r.body("Gaps whose clean signal is quieter than −60 dBFS RMS are excluded from this analysis, because an SNR "
           f"of near-silence is meaningless. Signal complexity is quantified by the mean spectral flatness of the "
           f"clean file (ratio of geometric to arithmetic mean of the power spectrum; 0 = tonal, 1 = noise-like) "
           f"{cite('gray1974')}.")

    r.heading("2.7  Statistical analysis", 2)
    r.body(f"The design is paired: both methods conceal the same files. For each metric, the per-file "
           f"difference is oriented so that a positive value always means 'the hybrid model is better'. "
           f"Differences are tested with the two-sided Wilcoxon signed-rank test {cite('wilcoxon1945')} (ties "
           f"dropped), which does not assume normality. Because four metrics are tested, p-values are adjusted "
           f"with the Holm-Bonferroni procedure {cite('holm1979')} and compared with α = 0.05. Effect size is "
           f"the matched-pairs rank-biserial correlation {cite('kerby2014')},")
    r.equation("r = (W⁺ − W⁻) / (W⁺ + W⁻)")
    r.body(f"where W⁺ and W⁻ are the rank sums of positive and negative differences (|r| < 0.1 negligible, "
           f"< 0.3 small, < 0.5 medium, otherwise large). The mean difference is reported with a 95% percentile "
           f"bootstrap confidence interval (10,000 resamples) {cite('efron1993')}. In the per-burst analysis, "
           f"bursts from the same file are not independent, so confidence intervals there use a cluster "
           f"bootstrap that resamples whole files.")

    r.heading("2.8  Hyperparameter sweeps", 2)
    rows = []
    for param, (df, _) in ctx.sweeps.items():
        label, key = SWEEP_LABELS.get(param, (param, param))
        values = sorted(df["value"].unique())
        rows.append([label, fg(ctx.hparams.get(key, np.nan)), ", ".join(fg(v) for v in values), df["file"].nunique()])
    r.body(f"Each hyperparameter was varied on its own while the others were held at their defaults "
           f"(one-at-a-time design). At every value, the same files were concealed whole with their real traces "
           f"by both methods and scored with all four metrics; compute time was recorded per file "
           f"(Table {r.next_tab()}). Burst length is deliberately not swept: in a causal, recursive concealer it "
           f"is a property of the network, not a design choice, and its effect is instead measured on the real "
           f"bursts (Section 3.4).")
    if rows:
        r.table(["Hyperparameter", "Default", "Values tested", "Files"], rows,
                "Hyperparameter sweep design.", col_widths_cm=[4.2, 1.8, 8.5, 1.5])

    r.heading("2.9  Computational cost", 2)
    info = ctx.run_info
    machine = (f"{info.get('processor') or info.get('platform', 'unknown CPU')}, {info.get('cpu_count', '?')} "
               f"logical cores, PyTorch using {info.get('torch_threads', '?')} threads") if info else "the evaluation machine"
    r.body(f"The wall-clock time of each method's concealment of each file was measured after a warm-up run "
           f"(model loading excluded). Since both methods only do work on lost packets, the time per file divided "
           f"by its number of lost packets gives the cost of concealing one packet. The real-time factor is")
    r.equation(f"RTF = t_packet / (P / f_s) = t_packet / {ctx.packet_ms:.1f} ms")
    r.body(f"A method can keep up with live audio only if RTF < 1. All timings were measured on CPU ({machine}).")


def section_results(r: Report, ctx: Context):
    r.heading("3  Results", 1)
    n = ctx.net

    # 3.1 ------------------------------------------------------------------
    r.heading("3.1  Characteristics of the real loss traces", 2)
    fig = r.next_fig()
    r.body(f"Figure {fig} shows the burst-length distribution and the per-file loss rate over all "
           f"{n['n_files']} traces. Losses are dominated by isolated packets ({n['pct_single']:.1f}% of bursts) "
           f"and {n['pct_le2']:.1f}% of all bursts are at most two packets ({2 * ctx.packet_ms:.1f} ms) long; "
           f"the longest burst spans {n['max_burst']} packets ({n['max_burst'] * ctx.packet_ms:.0f} ms), and "
           f"only {n['n_ge8']} bursts reach 8 packets or more. Typical concealment therefore has to bridge one or "
           f"two packets, but the rare long bursts are where errors accumulate in a recursive concealer.")
    r.figure(ctx.out / "burst_comparisons" / "network_statistics.png",
             "Loss statistics of the real packet traces. Left: number of loss bursts per burst length (log "
             "scale). Right: distribution of the packet loss rate across files; the dashed line marks the mean.")

    # 3.2 ------------------------------------------------------------------
    r.heading("3.2  Overall concealment quality (RQ1)", 2)
    if ctx.tests is None:
        r.para(("Significance tests not available -- run run_statistical_analysis.py.", "i"))
    else:
        tab = r.next_tab()
        r.body(f"Table {tab} summarizes the paired comparison over {ctx.n_files} files. "
               + significance_sentence(ctx))
        rows = []
        for _, t in ctx.tests.iterrows():
            rows.append([t["metric"], f(t["mean_lpc"], 3), f(t["mean_nn"], 3),
                         f"{fs(t['mean_nn_advantage'], 3)} [{fs(t['ci95_low'], 3)}, {fs(t['ci95_high'], 3)}]",
                         f"{t['nn_win_rate_pct']:.0f}%", fp(t["p_holm"]),
                         f"{t['rank_biserial_r']:+.2f} ({t['effect_size']})", t["verdict"]])
        r.table(["Metric", "AR (mean)", "Hybrid (mean)", "Hybrid advantage [95% CI]", "Hybrid wins",
                 "p (Holm)", "Effect size r", "Verdict"], rows,
                f"Paired comparison of AR-only and hybrid (AR + NN) concealment over {ctx.n_files} files.",
                col_widths_cm=[2.4, 1.5, 1.6, 3.3, 1.3, 1.4, 2.2, 2.3], font_size=8,
                notes="Advantage = hybrid minus AR, sign-flipped for MR-STFT so that positive always favours the "
                      "hybrid model. CI: bootstrap 95% confidence interval of the mean. p: two-sided Wilcoxon "
                      "signed-rank test, Holm-corrected over the four metrics. r: matched-pairs rank-biserial "
                      "correlation.")

        f_win, f_dist, f_diff = r.next_fig(), r.next_fig() + 1, r.next_fig() + 2
        r.body(f"The win rate (Figure {f_win}) only tells which method was better on each file. The per-file "
               f"distributions (Figure {f_dist}) and the paired differences (Figure {f_diff}) show how large the "
               f"differences are and how consistent they are across files.")
        r.figure(ctx.out / "metrics_comparison" / "win_rate_summary.png",
                 "Percentage of files on which the hybrid model scores better than AR alone, per metric. The "
                 "dashed line marks 50% (no systematic difference).", width_cm=14)
        r.figure(ctx.out / "statistical_analysis" / "metric_distributions.png",
                 "Per-file metric values of both methods (boxes: median and interquartile range; whiskers: 1.5 "
                 "IQR). Gray lines connect the two results of the same file.")
        r.figure(ctx.out / "statistical_analysis" / "paired_differences.png",
                 "Per-file difference between the methods, oriented so that values above zero favour the "
                 "hybrid model. Black marker: mean with bootstrap 95% confidence interval. Panel titles give "
                 "the Holm-corrected Wilcoxon p-value and the rank-biserial effect size.")

    # 3.3 ------------------------------------------------------------------
    r.heading("3.3  Dependence on signal complexity (RQ2)", 2)
    if ctx.per_file is not None:
        rows, notable = [], []
        for key, label, lib in ALL_METRICS:
            adv = oriented(ctx.per_file, key, lib)
            flat = ctx.per_file["complexity_flatness"]
            if adv.nunique() > 1 and flat.nunique() > 1:
                rho, p = stats.spearmanr(flat, adv)
            else:
                rho, p = np.nan, np.nan
            rows.append([label, fs(rho), fp(p)])
            if not np.isnan(p) and p < 0.05:
                notable.append(f"{label} (ρ = {rho:+.2f})")
        fig, tab = r.next_fig(), r.next_tab()
        text = (f"Figure {fig} plots each metric against the spectral flatness of the clean file, and Table {tab} "
                f"gives the Spearman rank correlation between flatness and the hybrid model's advantage. ")
        text += (f"The advantage correlated significantly with signal complexity for {join_words(notable)}."
                 if notable else
                 "No metric showed a significant correlation between signal complexity and the hybrid model's "
                 "advantage (uncorrected p ≥ 0.05).")
        r.body(text)
        r.figure(ctx.out / "metrics_comparison" / "complexity_grid.png",
                 "Concealment quality versus signal complexity (mean spectral flatness of the clean file, log "
                 "scale). One dot per file and method; dashed lines are least-squares trends in log-flatness.")
        r.table(["Metric", "Spearman ρ (flatness vs. hybrid advantage)", "p (uncorrected)"], rows,
                "Correlation between signal complexity and the hybrid model's advantage.",
                col_widths_cm=[4, 8, 4])

    # 3.4 ------------------------------------------------------------------
    r.heading("3.4  Dependence on burst length (RQ2)", 2)
    if ctx.burst_q is None or ctx.per_burst is None:
        r.para(("Per-burst results not available -- re-run run_metrics_comparison.py and "
                "run_statistical_analysis.py.", "i"))
    else:
        audible = ctx.per_burst[ctx.per_burst["clean_gap_rms_db"] > -60]
        fig, tab = r.next_fig(), r.next_tab()
        r.body(f"Every real burst in the evaluated files was scored with the gap SNR: {len(audible):,} bursts from "
               f"{audible['file'].nunique()} files after excluding {len(ctx.per_burst) - len(audible):,} "
               f"near-silent gaps (Figure {fig}, Table {tab}). " + burst_sentence(ctx))
        r.figure(ctx.out / "statistical_analysis" / "burst_length_quality.png",
                 "Gap SNR versus real burst length. Left: mean gap SNR of each method. Right: mean difference "
                 "(hybrid minus AR). Error bars: 95% confidence intervals from a cluster bootstrap over files; "
                 "n is the number of bursts per bin.")
        rows = [[b["burst_length_bin"], int(b["n_bursts"]), int(b["n_files"]),
                 f"{f(b['gap_snr_lpc_mean'])} [{f(b['gap_snr_lpc_ci_low'])}, {f(b['gap_snr_lpc_ci_high'])}]",
                 f"{f(b['gap_snr_nn_mean'])} [{f(b['gap_snr_nn_ci_low'])}, {f(b['gap_snr_nn_ci_high'])}]",
                 f"{fs(b['nn_advantage_mean'])} [{fs(b['nn_advantage_ci_low'])}, {fs(b['nn_advantage_ci_high'])}]",
                 f"{b['nn_win_rate_pct']:.0f}%"] for _, b in ctx.burst_q.iterrows()]
        r.table(["Burst (packets)", "Bursts", "Files", "AR gap SNR (dB)", "Hybrid gap SNR (dB)",
                 "Hybrid advantage (dB)", "Hybrid wins"], rows,
                "Gap SNR by burst length (mean [95% cluster-bootstrap CI]).",
                col_widths_cm=[1.8, 1.4, 1.2, 3.2, 3.2, 3.4, 1.8], font_size=8)

    # 3.5 ------------------------------------------------------------------
    r.heading("3.5  Hyperparameter sensitivity (RQ3)", 2)
    if not ctx.sweeps:
        r.para(("Sweep results not available -- run run_hparam_sweeps.py.", "i"))
    else:
        first = r.next_fig()
        r.body(f"Figures {first}-{first + len(ctx.sweeps) - 1} show every metric and the compute time as a "
               f"function of each hyperparameter, and Table {r.next_tab()} lists the best value found for SNR and "
               f"PLCMOS. Because the network was trained on the residual of the default AR configuration, values "
               f"away from the default change both the AR prediction and the input statistics the network sees; "
               f"the hybrid results at non-default settings therefore measure robustness to this mismatch, not "
               f"the performance of a network retrained for that setting.")
        for param, (df, png) in ctx.sweeps.items():
            label, key = SWEEP_LABELS.get(param, (param, param))
            by = df.groupby("value")[["snr_lpc", "snr_nn", "plcmos_lpc", "plcmos_nn"]].mean()
            nn_snr = int((by["snr_nn"] > by["snr_lpc"]).sum())
            nn_mos = int((by["plcmos_nn"] > by["plcmos_lpc"]).sum())
            n_sw = df["file"].nunique()
            files_txt = f"{n_sw} file{'s' if n_sw != 1 else ''} per value"
            r.heading(f"3.5.{list(ctx.sweeps).index(param) + 1}  {label}", 3)
            r.body(f"Over the {len(by)} values tested ({files_txt}), the hybrid model "
                   f"had the higher mean SNR at {nn_snr} and the higher mean PLCMOS at {nn_mos} of them. Mean SNR "
                   f"ranged from {f(by['snr_lpc'].min())} to {f(by['snr_lpc'].max())} dB for AR and from "
                   f"{f(by['snr_nn'].min())} to {f(by['snr_nn'].max())} dB for the hybrid model.")
            r.figure(png, f"Sweep of the {label} (default {fg(ctx.hparams.get(key, np.nan))}; other "
                          f"hyperparameters at their defaults; {files_txt}). Panels: "
                          f"SNR, MR-STFT distance, PLCMOS and ViSQOL (mean ± 1 SD across files) and compute time "
                          f"per file (median and interquartile range).")
        if ctx.sweep_summary is not None:
            ss = ctx.sweep_summary
            rows = []
            for (param, method), g in ss.groupby(["hyperparameter", "method"], sort=False):
                label = SWEEP_LABELS.get(param, (param,))[0]
                cells = [label, "Hybrid" if method == "NN" else "AR"]
                for metric in ("SNR (dB)", "PLCMOS"):
                    row = g[g["metric"] == metric]
                    if row.empty:
                        cells.append("n/a")
                        continue
                    row = row.iloc[0]
                    cells.append(f"{fg(row['best_value'])} ({fs(row['gain_over_default'], 3)})")
                row = g.iloc[0]
                cells.append(f"{f(row['median_time_ms_at_default'], 0)} → {f(row['median_time_ms_at_best'], 0)}")
                rows.append(cells)
            r.table(["Hyperparameter", "Method", "Best for SNR (gain)", "Best for PLCMOS (gain)",
                     "Time/file (ms), default → best"], rows,
                    "Best value of each hyperparameter and gain over the default (positive = better than default).",
                    col_widths_cm=[4, 1.6, 3.4, 3.4, 3.6], font_size=8,
                    notes="Gains are differences of means across files and are not tested for significance. "
                          "Time is the median per-file concealment time at the best SNR value. The full table "
                          "for all four metrics is in Appendix C.")

    # 3.6 ------------------------------------------------------------------
    r.heading("3.6  Computational cost (RQ4)", 2)
    if not ctx.has_timing:
        r.para(("Timing data not available -- re-run run_metrics_comparison.py.", "i"))
    else:
        rows = []
        for method, name in (("lpc", "AR only"), ("nn", "Hybrid (AR + NN)")):
            ms = ctx.per_file[f"ms_per_lost_packet_{method}"].dropna()
            rtf = ctx.rtf(method)
            rows.append([name, f"{np.median(ms):.2f}", f"{np.percentile(ms, 25):.2f}-{np.percentile(ms, 75):.2f}",
                         f"{np.median(rtf):.3f}", "yes" if np.median(rtf) < 1 else "no"])
        fig, tab = r.next_fig(), r.next_tab()
        ratio = np.median(ctx.per_file["ms_per_lost_packet_nn"].dropna()) / np.median(
            ctx.per_file["ms_per_lost_packet_lpc"].dropna())
        r.body(f"Table {tab} and Figure {fig} report the cost of concealing one packet. "
               + realtime_sentence(ctx) + f" The neural residual multiplies the per-packet cost by a factor of "
               f"about {ratio:.0f}. The challenge rules do not disqualify slower-than-real-time systems but encourage "
               f"real-time operation {cite('challenge2024')}.")
        r.table(["Method", "Median ms / packet", "IQR (ms)", "Median RTF", "Real time?"], rows,
                f"Per-packet concealment cost (packet duration {ctx.packet_ms:.1f} ms).",
                col_widths_cm=[4, 3, 3, 3, 3])
        r.figure(ctx.out / "statistical_analysis" / "realtime_feasibility.png",
                 "Real-time factor (time to conceal one packet divided by the packet duration) per file, log "
                 "scale. Below the red line a method keeps up with live audio on this machine.", width_cm=13)

    # 3.7 ------------------------------------------------------------------
    r.heading("3.7  Qualitative examples", 2)
    sets = burst_figure_sets(ctx.out / "burst_comparisons")
    if not sets:
        r.para(("Burst examples not available -- run plot_burst_comparison.py.", "i"))
    else:
        r.body("To show what the metrics above summarize, one real burst of each length was selected and "
               "plotted in three ways: the waveform and spectrogram around the gap (shaded) for both methods; the "
               "sample-wise difference between the two concealments; and a zoom on the hand-back point where "
               "concealment crossfades into the next received packet, together with the clean reference. Table "
               f"{r.next_tab()} lists the mean squared difference between the two methods inside each gap.")
        if ctx.gap_mse is not None:
            col = [c for c in ctx.gap_mse.columns if "_vs_" in c][0]
            r.table(["Burst (packets)", "File", "Start packet", "Gap MSE between methods"],
                    [[int(g["burst_length"]), g["file"][:28] + "…", int(g["start_packet"]), f"{g[col]:.3e}"]
                     for _, g in ctx.gap_mse.sort_values("burst_length").iterrows()],
                    "Selected example bursts.", col_widths_cm=[2.5, 7, 2.5, 4])
        for i, (length, tag, paths) in enumerate(sets, start=1):
            r.heading(f"3.7.{i}  {length}-packet burst ({length * ctx.packet_ms:.1f} ms)", 3)
            r.figure(paths["compare"], f"{length}-packet burst in file {tag}: waveform (top) and spectrogram "
                                       f"(bottom) of the AR (left) and hybrid (right) concealment. Shaded: lost "
                                       f"samples.")
            r.figure(paths["diff"], f"{length}-packet burst: AR concealment minus hybrid concealment; the "
                                    f"gap-MSE in the title quantifies the neural contribution.", width_cm=11)
            r.figure(paths["crossfade"], f"{length}-packet burst: zoom on the hand-back point (dashed line) where "
                                         f"concealment ends; the shaded region is the {ctx.hparams['extra_dim']}-"
                                         f"sample crossfade into the received packet.")

    # 3.8 ------------------------------------------------------------------
    r.heading("3.8  Failure cases", 2)
    if ctx.failures is None or ctx.failures.empty:
        r.para(("No failure cases available (either the hybrid model never lost on the selected metric, or "
                "run_statistical_analysis.py has not been run).", "i"))
    else:
        metric_key = [c[:-4] for c in ctx.failures.columns if c.endswith("_lpc")][0]
        label = next(l for k, l, _ in ALL_METRICS if k == metric_key)
        flat_med = ctx.per_file["complexity_flatness"].median()
        loss_med = ctx.per_file["loss_rate_pct"].median()
        r.body(f"Reporting only favourable examples would bias the picture. The {len(ctx.failures)} files on "
               f"which AR alone beat the hybrid model by the largest {label} margin are listed in Table "
               f"{r.next_tab()}; for each, the burst with the largest AR advantage in gap SNR is plotted. Their "
               f"spectral flatness ranged from {ctx.failures['complexity_flatness'].min():.2e} to "
               f"{ctx.failures['complexity_flatness'].max():.2e} (median of all evaluated files: {flat_med:.2e}), "
               f"and their loss rates from {ctx.failures['loss_rate_pct'].min():.1f}% to "
               f"{ctx.failures['loss_rate_pct'].max():.1f}% (median {loss_med:.1f}%).")
        r.table(["Case", "File", f"AR {label}", f"Hybrid {label}", "Advantage", "Flatness", "Loss rate",
                 "Plotted burst"],
                [[int(c["case"]), c["file"][:22] + "…", f(c[f"{metric_key}_lpc"]), f(c[f"{metric_key}_nn"]),
                  fs(c["nn_advantage"]), f"{c['complexity_flatness']:.2e}", f"{c['loss_rate_pct']:.1f}%",
                  f"{int(c['plotted_burst_length'])} pkt @ {int(c['plotted_burst_start'])}"]
                 for _, c in ctx.failures.iterrows()],
                f"Files with the largest AR advantage on {label}.",
                col_widths_cm=[1, 4, 1.8, 1.8, 1.8, 1.8, 1.6, 2.2], font_size=8)
        fail_dir = ctx.out / "statistical_analysis" / "failure_cases"
        for length, tag, paths in burst_figure_sets(fail_dir, "failure"):
            case = tag.split("_", 1)[0].replace("failure", "")
            r.heading(f"3.8.{case}  Failure case {case}", 3)
            r.figure(paths["compare"], f"Failure case {case}: waveform and spectrogram around the plotted "
                                       f"{length}-packet burst, AR (left) and hybrid (right).")
            r.figure(paths["diff"], f"Failure case {case}: AR minus hybrid concealment.", width_cm=11)
            r.figure(paths["crossfade"], f"Failure case {case}: hand-back crossfade with the clean reference.")


def section_discussion(r: Report, ctx: Context):
    r.heading("4  Discussion", 1)
    if ctx.tests is not None:
        adv = {k: ctx.test_row(k)["mean_nn_advantage"] for k, _, _ in ALL_METRICS}
        sig = {k: ctx.test_row(k)["verdict"] for k, _, _ in ALL_METRICS}
        spectral_pos = [METRIC_DESC[k] for k in ("mrstft", "plcmos", "visqol") if adv[k] > 0]
        r.heading("4.1  What the neural residual adds", 2)
        if adv["snr"] <= 0 and spectral_pos:
            r.body(f"The two metric families point in different directions. On average the hybrid model "
                   f"{'lowered' if adv['snr'] < 0 else 'did not change'} the waveform SNR "
                   f"({fs(adv['snr'])} dB, {sig['snr']}) while it improved {join_words(spectral_pos)}. "
                   f"This pattern is consistent with how the network is trained: its loss includes a spectral "
                   f"term, and a spectral improvement can coexist with a larger sample-wise error, because a "
                   f"prediction whose spectrum is right but whose phase drifts is penalized by SNR and not by "
                   f"magnitude-based or perceptual measures. One interpretation is that the residual makes the "
                   f"concealment more plausible rather than more exact; testing this directly would require a "
                   f"listening test (Section 5).")
        elif all(v > 0 for v in adv.values()):
            r.body("All four metrics favoured the hybrid model on average, so the neural residual improved both "
                   "the waveform match and the perceptual plausibility of the concealment.")
        else:
            r.body("The metrics did not favour either method consistently: "
                   + "; ".join(f"{METRIC_DESC[k]} {fs(v, 3)} ({sig[k]})" for k, v in adv.items()) + ".")
        if ctx.preliminary:
            r.body(f"With only {ctx.n_files} files, none of these directions can be regarded as established; "
                   f"they should be read as hypotheses for the full experiment.")

    r.heading("4.2  Burst length, complexity and recursion", 2)
    observed = ""
    if ctx.burst_q is not None and len(ctx.burst_q) > 1:
        first, last = ctx.burst_q.iloc[0], ctx.burst_q.iloc[-1]
        fell = (last["gap_snr_lpc_mean"] < first["gap_snr_lpc_mean"] and
                last["gap_snr_nn_mean"] < first["gap_snr_nn_mean"])
        observed = (", as observed here for both methods" if fell else
                    ", although this was not observed consistently in the present data")
    r.body("Because both concealers are recursive, every packet of a burst is predicted from the previous "
           f"predictions, so errors accumulate and gap quality is expected to fall with burst length{observed}. "
           "Whether the neural residual helps more on long bursts is the practically important "
           f"question, but long bursts are rare in real traces (Table {ctx.net_table}), which makes the long-burst bins the "
           "least certain ones; their confidence intervals should be read accordingly.")

    r.heading("4.3  Cost-quality trade-off", 2)
    if ctx.has_timing:
        nn_rtf, lpc_rtf = np.median(ctx.rtf("nn")), np.median(ctx.rtf("lpc"))
        r.body(f"The AR concealer {'runs in real time' if lpc_rtf < 1 else 'does not run in real time'} on the "
               f"evaluation CPU (median RTF {lpc_rtf:.2f}), whereas the hybrid "
               f"model {'exceeded' if nn_rtf > 1 else 'stayed within'} the real-time budget "
               f"(median RTF {nn_rtf:.2f}). Any quality benefit of the residual must therefore be weighed against "
               f"a cost about {nn_rtf / lpc_rtf:.0f} times higher; options such as multi-threaded or GPU inference, model "
               f"quantization or a smaller network would need to be considered for live deployment.")
    r.heading("4.4  Coupling between the AR configuration and the network", 2)
    r.body("The network was trained on residuals of the default AR configuration. Changing the AR order, "
           "context length or regularization changes those residuals, so the hybrid model operates outside its "
           "training conditions. The sweeps therefore measure how robust a fixed network is to such changes; a "
           "fair test of a different AR configuration for the hybrid model would require retraining the network "
           "for it.")


def section_limitations(r: Report, ctx: Context):
    r.heading("5  Limitations and future work", 1)
    if ctx.preliminary:
        sweep_n = sorted(set(ctx.sweep_files.values()))
        sweep_txt = (f" and on sweeps with {join_words([str(v) for v in sweep_n])} "
                     f"file{'s' if sweep_n != [1] else ''} per value") if sweep_n else ""
        r.bullet(("Sample size. ", "b"), f"This version is based on {ctx.n_files} files{sweep_txt}, too few "
                 f"for reliable conclusions. The full runs will replace these numbers.")
    r.bullet(("No listening test. ", "b"), f"The challenge itself evaluates submissions with a MUSHRA listening "
             f"test {cite('itu1534')}, noting that no objective metric has been shown to correlate reliably with "
             f"perception for music PLC {cite('challenge2024')}. PLCMOS was designed for speech, so the "
             f"perceptual metrics here are proxies; a MUSHRA test is the natural next step.")
    r.bullet(("One dataset, one checkpoint. ", "b"), "All results come from one music test set and the "
             "pretrained challenge baseline. Generalization to other content (for example the speech validation "
             "set of the ICASSP 2024 PLC Challenge, which uses 48 kHz audio and 20 ms packets) and to retrained "
             "networks remains to be tested.")
    r.bullet(("Gap SNR is phase-sensitive. ", "b"), "The per-burst analysis uses SNR because the perceptual "
             "models need longer inputs; it penalizes plausible but phase-shifted concealments.")
    r.bullet(("One-at-a-time sweeps. ", "b"), "Interactions between hyperparameters (for example AR order and "
             "context length) are not explored; a joint search would be needed to find a global optimum.")
    r.bullet(("Timing on one CPU. ", "b"), "Real-time factors depend on hardware, threading and implementation; "
             "they are indicative of the relative cost, not of every deployment.")


def section_conclusion(r: Report, ctx: Context):
    r.heading("6  Conclusion", 1)
    text = (f"This report compared AR-only and hybrid AR + neural (PARCnet) packet loss concealment for "
            f"networked music on real loss traces, using two independently implemented concealers that are "
            f"identical except for the neural residual. ")
    text += " ".join(s for s in [significance_sentence(ctx), burst_sentence(ctx), realtime_sentence(ctx)] if s)
    if ctx.preliminary:
        text += " These conclusions are preliminary and will be updated with the full experiments."
    r.body(text)


def section_references(r: Report):
    r.heading("References", 1)
    for i, (_, text) in enumerate(REFERENCES, start=1):
        p = r.para(f"[{i}] {text}", size=9.5, space_after=3)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)


def section_appendix(r: Report, ctx: Context):
    r.page_break()
    r.heading("Appendix A  Reproducibility", 1)
    r.body("All results were produced with the scripts in the research/ folder of the project repository, in "
           "this order:")
    for cmd in [f"python -m research.run_metrics_comparison --n-files {ctx.n_files} --seed "
                f"{ctx.run_info.get('seed', 0)}",
                f"python -m research.run_hparam_sweeps --n-files {max(ctx.sweep_files.values(), default=20)}",
                "python -m research.plot_burst_comparison",
                "python -m research.run_statistical_analysis",
                "python -m research.build_report"]:
        p = r.para(cmd, size=9)
        p.runs[0].font.name = "Consolas"
        p.paragraph_format.left_indent = Cm(0.8)
    info = ctx.run_info
    rows = [["Platform", info.get("platform", "n/a")], ["Processor", info.get("processor", "n/a")],
            ["Logical cores", info.get("cpu_count", "n/a")], ["PyTorch threads", info.get("torch_threads", "n/a")]]
    rows += software_versions()
    r.table(["Item", "Value"], rows, "Hardware of the main evaluation run and software versions.",
            col_widths_cm=[5, 11])

    if ctx.per_file is not None:
        r.heading("Appendix B  Descriptive statistics", 1)
        rows = []
        for key, label, _ in ALL_METRICS:
            for method, name in (("lpc", "AR"), ("nn", "Hybrid")):
                v = ctx.per_file[f"{key}_{method}"]
                rows.append([label, name, f(v.mean(), 3), f(v.std(), 3), f(v.median(), 3),
                             f"{v.quantile(0.25):.3f}-{v.quantile(0.75):.3f}", f"{v.min():.3f}-{v.max():.3f}"])
        r.table(["Metric", "Method", "Mean", "SD", "Median", "IQR", "Range"], rows,
                f"Descriptive statistics of all metrics over {ctx.n_files} files.",
                col_widths_cm=[3, 1.8, 1.8, 1.8, 1.8, 2.9, 2.9], font_size=8)

    if ctx.sweep_summary is not None:
        r.heading("Appendix C  Full hyperparameter sweep summary", 1)
        ss = ctx.sweep_summary
        rows = [[SWEEP_LABELS.get(s["hyperparameter"], (s["hyperparameter"],))[0],
                 "Hybrid" if s["method"] == "NN" else "AR", s["metric"], fg(s["best_value"]),
                 f(s["metric_at_default"], 3), f(s["metric_at_best"], 3), fs(s["gain_over_default"], 3)]
                for _, s in ss.iterrows()]
        r.table(["Hyperparameter", "Method", "Metric", "Best value", "At default", "At best", "Gain"], rows,
                "Best value per hyperparameter, method and metric.",
                col_widths_cm=[3.6, 1.5, 2.8, 1.8, 2.1, 2.1, 2.1], font_size=7.5)


def build(output_dir: Path, out_path: Path) -> Path:
    ctx = Context(output_dir)
    r = Report()
    section_front(r, ctx)
    section_introduction(r, ctx)
    section_methods(r, ctx)
    section_results(r, ctx)
    section_discussion(r, ctx)
    section_limitations(r, ctx)
    section_conclusion(r, ctx)
    section_references(r)
    section_appendix(r, ctx)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    return safe_io._save_with_fallback(lambda p: r.doc.save(str(p)), out_path, "_locked")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR,
                        help="Folder containing the four scripts' output subfolders.")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="Path of the .docx to write.")
    args = parser.parse_args()
    path = build(args.output_dir, args.out)
    print(f"Report saved to {path}")


if __name__ == "__main__":
    main()
